import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generate_excel_report(title, columns, data):
    """
    Genera un archivo Excel (.xlsx) en memoria.
    :param title: Título de la hoja
    :param columns: Lista de nombres de columnas o tuplas (header_label, data_key)
    :param data: Lista de diccionarios con los datos
    :return: io.BytesIO
    """
    wb = Workbook()
    ws = wb.active
    ws.title = str(title)[:31]

    # Normalizar columnas a pares (header_label, key)
    col_pairs = []
    for col in columns:
        if isinstance(col, tuple):
            col_pairs.append(col)
        elif isinstance(col, dict):
            col_pairs.append((col.get('header', col.get('key')), col.get('key')))
        else:
            col_pairs.append((str(col).replace('_', ' ').title(), str(col)))

    # Encabezados
    headers = [p[0] for p in col_pairs]
    ws.append(headers)

    # Estilos encabezado
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="F59E0B")
    thin_border = Border(
        left=Side(style='thin', color='334155'),
        right=Side(style='thin', color='334155'),
        top=Side(style='thin', color='334155'),
        bottom=Side(style='thin', color='334155')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 24

    # Filas de datos
    for row in data:
        row_vals = []
        for _, key in col_pairs:
            val = row.get(key, '')
            if val is None:
                val = ''
            row_vals.append(val)
        ws.append(row_vals)

    # Auto-ajustar ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_csv_report(title, columns, data):
    """
    Genera un archivo CSV en memoria como BytesIO.
    :param title: Título del reporte (no usado en CSV, para mantener signatura)
    :param columns: Lista de columnas
    :param data: Lista de diccionarios con datos
    :return: io.BytesIO
    """
    col_pairs = []
    for col in columns:
        if isinstance(col, tuple):
            col_pairs.append(col)
        elif isinstance(col, dict):
            col_pairs.append((col.get('header', col.get('key')), col.get('key')))
        else:
            col_pairs.append((str(col).replace('_', ' ').title(), str(col)))

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

    # Header
    writer.writerow([p[0] for p in col_pairs])

    # Data
    for row in data:
        row_vals = []
        for _, key in col_pairs:
            val = row.get(key, '')
            if val is None:
                val = ''
            row_vals.append(str(val))
        writer.writerow(row_vals)

    bytes_output = io.BytesIO(output.getvalue().encode('utf-8-sig'))
    bytes_output.seek(0)
    return bytes_output
